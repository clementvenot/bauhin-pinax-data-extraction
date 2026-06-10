from difflib import SequenceMatcher
import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook

REFERENCE_XLSX_PATH = Path(__file__).with_name("references.xlsx")


def _normalize_reference_text(text):
    if text is None:
        return ""

    normalized = unicodedata.normalize("NFKD", str(text))
    normalized = "".join(character for character in normalized if not unicodedata.combining(character))
    normalized = normalized.lower()
    normalized = re.sub(r"[^\w]+", " ", normalized, flags=re.UNICODE)
    return re.sub(r"\s+", " ", normalized).strip()


def _tokenize_with_spans(text):
    tokens = []
    for match in re.finditer(r"\w+", text or "", flags=re.UNICODE):
        tokens.append(
            {
                "text": match.group(0),
                "norm": _normalize_reference_text(match.group(0)),
                "start": match.start(),
                "end": match.end(),
            }
        )
    return tokens


def _resolve_reference_id(ref_value):
    if isinstance(ref_value, (list, tuple, set)):
        return max(int(value) for value in ref_value)
    return int(ref_value)


def load_reference_from_xlsx(xlsx_path, sheet_name=None):
    """Load a reference dictionary from an Excel file.

    The file must contain two columns:
    - column 1: id
    - column 2: reference text

    The first row is ignored.
    Returns a dictionary shaped as {reference_text: id}.
    If a reference text appears multiple times, the highest id is kept.
    """

    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    worksheet = workbook[sheet_name] if sheet_name else workbook.active

    loaded_reference = {}
    for row in worksheet.iter_rows(min_row=2, min_col=1, max_col=2, values_only=True):
        reference_id, reference_text = row
        if reference_id is None or reference_text is None:
            continue

        cleaned_text = str(reference_text).strip()
        if not cleaned_text:
            continue

        current_id = int(reference_id)
        previous_id = loaded_reference.get(cleaned_text)
        if previous_id is None or current_id > previous_id:
            loaded_reference[cleaned_text] = current_id

    workbook.close()
    return loaded_reference


try:
    reference = load_reference_from_xlsx(REFERENCE_XLSX_PATH)
except Exception:
    reference = {}


def _expand_removal_span(text, start, end):
    expanded_start = start
    expanded_end = end

    while expanded_start > 0 and text[expanded_start - 1].isspace():
        expanded_start -= 1

    while expanded_start > 0 and not text[expanded_start - 1].isalnum() and text[expanded_start - 1] != "_":
        expanded_start -= 1

    while expanded_end < len(text) and not text[expanded_end].isalnum() and text[expanded_end] != "_":
        expanded_end += 1

    while expanded_end < len(text) and text[expanded_end].isspace():
        expanded_end += 1

    return expanded_start, expanded_end


def find_fuzzy_references(text, min_score=0.82):
    """Return matched reference ids followed by the remaining unmatched text.

    The returned format is:
    "id1,id2, remaining text"
    """

    source_text = text or ""
    tokens = _tokenize_with_spans(source_text)
    if not tokens:
        return ""

    candidates = []

    for ref_text, ref_value in reference.items():
        normalized_reference = _normalize_reference_text(ref_text)
        if not normalized_reference:
            continue

        ref_id = _resolve_reference_id(ref_value)

        reference_tokens = normalized_reference.split()
        token_count = len(reference_tokens)
        window_min = max(1, token_count - 1)
        window_max = token_count + 1

        for start_index in range(len(tokens)):
            for window_size in range(window_min, window_max + 1):
                end_index = start_index + window_size
                if end_index > len(tokens):
                    break

                candidate_text = " ".join(token["norm"] for token in tokens[start_index:end_index] if token["norm"])
                if not candidate_text:
                    continue

                score = SequenceMatcher(None, normalized_reference, candidate_text).ratio()
                if score >= min_score:
                    candidates.append(
                        {
                            "score": score,
                            "start": start_index,
                            "end": end_index,
                            "ref_id": ref_id,
                        }
                    )

    candidates.sort(key=lambda candidate: (-candidate["score"], candidate["start"], -(candidate["end"] - candidate["start"])))

    used_token_indexes = set()
    selected = []

    for candidate in candidates:
        candidate_range = set(range(candidate["start"], candidate["end"]))
        if candidate_range & used_token_indexes:
            continue

        used_token_indexes.update(candidate_range)
        selected.append(candidate)

    selected.sort(key=lambda candidate: candidate["start"])

    matched_ids = []
    seen_ids = set()
    for candidate in selected:
        ref_id = candidate["ref_id"]
        if ref_id not in seen_ids:
            seen_ids.add(ref_id)
            matched_ids.append(str(ref_id))

    if not matched_ids:
        return source_text.strip()

    removed_ranges = []
    for candidate in selected:
        start_token = tokens[candidate["start"]]
        end_token = tokens[candidate["end"] - 1]
        removed_ranges.append(_expand_removal_span(source_text, start_token["start"], end_token["end"]))

    removed_ranges.sort()
    merged_ranges = []
    for start, end in removed_ranges:
        if not merged_ranges or start > merged_ranges[-1][1]:
            merged_ranges.append([start, end])
        else:
            merged_ranges[-1][1] = max(merged_ranges[-1][1], end)

    remaining_parts = []
    cursor = 0
    for start, end in merged_ranges:
        remaining_parts.append(source_text[cursor:start])
        cursor = end
    remaining_parts.append(source_text[cursor:])

    remaining_text = []
    for part in remaining_parts:
        if not part:
            continue
        if remaining_text and remaining_text[-1] and remaining_text[-1][-1].isalnum() and part[0].isalnum():
            remaining_text.append(" ")
        remaining_text.append(part)

    remaining_text = re.sub(r"\s+", " ", "".join(remaining_text)).strip(" ,;:\t\n\r")

    if remaining_text:
        return f"{','.join(matched_ids)}, {remaining_text}"

    return ",".join(matched_ids)

# print dico from xlsx for debugging

if __name__ == "__main__":
    print(reference)
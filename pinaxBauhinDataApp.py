import json
import queue
import threading
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from pathlib import Path

from PIL import Image, ImageOps, ImageTk
from openpyxl import Workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
import pypdfium2 as pdfium

from remove_word_fuzzy import remove_fuzzy_phrase
from ocr_bicolumns import extract_line_indentation, delete_bars, pytesseract
from pipeline import pipeline as detect_page_regions
import reference_bauhin_fuzzy as reference_fuzzy

RECT_COLORS = {
    "titles": "#2d7ff9",
    "localisation": "#ffd900",
    "bicolumn_text": "#2ca02c",
    "italic_text": "#d62728",
}

RECT_STROKE_WIDTH = 3

RECT_LABELS_TO_DISPLAY = {
    "titles": "genre",
    "localisation": "localisation",
    "italic_text": "italique",
    "bicolumn_text": "bi-colonne",
}

RECT_LABELS_TO_INTERNAL = {display: internal for internal, display in RECT_LABELS_TO_DISPLAY.items()}


def clamp(value, minimum, maximum):
    return max(minimum, min(maximum, value))


def label_to_display(label):
    return RECT_LABELS_TO_DISPLAY.get(label, label)


def label_to_internal(label):
    return RECT_LABELS_TO_INTERNAL.get(label, label)


def rect_color(label):
    return RECT_COLORS.get(label_to_internal(label), RECT_COLORS["titles"])
def is_italic_name(normalized_name):
    return "italic" in normalized_name or "italique" in normalized_name


def replace_first_taxon_punctuation(text):
    text = text or ""
    for index, character in enumerate(text):
        if character in {".", ","}:
            return f"{text[:index]}%{text[index + 1:]}"
    return text


class PinaxBauhinDataApp(tk.Tk):
    def __init__(self):
        super().__init__()

        self.title("Caspar Bauhin - Extraction de données")
        self.geometry("1920x1080")

        self.pdf_path = ""
        self.document = None
        self.page_count = 0
        self.current_page = 0
        self.render_scale = 4.17

        self.page_annotations = []
        self.page_cache = {}
        self.loading_pages = set()
        self.load_queue = queue.Queue()
        self._poll_after_id = None

        self.current_image = None
        self.current_photo = None
        self.source_image_size = None
        self.display_image_size = None
        self.current_page_data = []

        self.canvas_image_id = None
        self.rect_items = {}
        self.handle_items = {}
        self.delete_cross_items = {}
        self.separator_items = {}
        self.temp_rect_id = None
        self.active_rect_uid = None
        self.action_mode = None
        self.drag_start_image = None
        self.drag_anchor_rect = None
        self.new_rect_label = tk.StringVar(value=label_to_display("titles"))
        self._next_rect_uid = 1

        self.gallery_window = None
        self.gallery_export_items = []
        self.validated_titles = []
        self.validated_italics = []
        self._validated_title_counter = 1
        self._last_validated_border_page_index = None
        self.reference_dictionary = None
        self.reference_dictionary_path = None
        self.export_button_text = tk.StringVar(value="Importer reference")

        self.view_scale = 1.0
        self.offset_x = 0
        self.offset_y = 0
        self.canvas_placeholder_text = "Aucune page à afficher"

        self._build_ui()
        self.bind("<Delete>", self._delete_selected_rectangle)
        self.bind("<Escape>", self._cancel_action)

        self.after(100, self._poll_load_queue)
        self.after(100, self._prompt_for_pdf)

    def _build_ui(self):
        self.columnconfigure(0, weight=1)
        self.rowconfigure(0, weight=1)

        main = ttk.Frame(self, padding=8)
        main.grid(row=0, column=0, sticky="nsew")
        main.columnconfigure(0, weight=1)
        main.columnconfigure(1, weight=7)
        main.rowconfigure(1, weight=1)
        main.rowconfigure(2, weight=0)

        toolbar = ttk.Frame(main)
        toolbar.grid(row=0, column=0, sticky="ew", pady=(0, 8))
        toolbar.columnconfigure(0, weight=0)
        toolbar.columnconfigure(1, weight=1)

        toolbar_left = ttk.Frame(toolbar)
        toolbar_left.grid(row=0, column=0, sticky="w")

        toolbar_right = ttk.Frame(toolbar)
        toolbar_right.grid(row=0, column=1, sticky="e")

        ttk.Button(toolbar_left, text="Ouvrir PDF", command=self._prompt_for_pdf).grid(row=0, column=0, padx=(0, 6))

        ttk.Button(toolbar_right, text="<", command=self.prev_page).grid(row=0, column=0, padx=(0, 6))
        ttk.Button(toolbar_right, text=">", command=self.next_page).grid(row=0, column=1, padx=(0, 12))

        self.page_label = ttk.Label(toolbar_right, text="Page 0 / 0")
        self.page_label.grid(row=0, column=2, padx=(0, 12))

        ttk.Label(toolbar_right, text="Aller à la page:").grid(row=0, column=3, padx=(0, 4))
        self.page_jump_var = tk.StringVar(value="1")
        self.page_jump_entry = ttk.Entry(toolbar_right, textvariable=self.page_jump_var, width=8, justify="right")
        self.page_jump_entry.grid(row=0, column=4, padx=(0, 4))
        self.page_jump_entry.bind("<Return>", self.go_to_page_from_entry)

        canvas_frame = ttk.Frame(main, borderwidth=1, relief="solid")
        canvas_frame.grid(row=1, column=0, sticky="nsew")
        canvas_frame.rowconfigure(0, weight=1)
        canvas_frame.columnconfigure(0, weight=1)

        self.canvas = tk.Canvas(canvas_frame, bg="#1f2328", highlightthickness=0)
        self.canvas.grid(row=0, column=0, sticky="nsew")

        self.canvas.bind("<Configure>", self._on_canvas_resize)
        self.canvas.bind("<ButtonPress-1>", self._on_canvas_press)
        self.canvas.bind("<B1-Motion>", self._on_canvas_drag)
        self.canvas.bind("<ButtonRelease-1>", self._on_canvas_release)

        right_panel = ttk.Frame(main, padding=(12, 0, 0, 0))
        right_panel.grid(row=0, column=1, rowspan=3, sticky="nsew")
        right_panel.columnconfigure(0, weight=1)
        right_panel.rowconfigure(1, weight=1)
        right_panel.rowconfigure(3, weight=1)

        ttk.Label(right_panel, text="Genres validés", font=("Helvetica", 13, "bold")).grid(row=0, column=0, sticky="w", pady=(0, 8))

        titles_container = ttk.Frame(right_panel)
        titles_container.grid(row=1, column=0, sticky="nsew")
        titles_container.columnconfigure(0, weight=1)
        titles_container.rowconfigure(0, weight=1)

        self.validated_titles_canvas = tk.Canvas(titles_container, highlightthickness=0, bg="#ffffff")
        self.validated_titles_scrollbar = ttk.Scrollbar(titles_container, orient="vertical", command=self.validated_titles_canvas.yview)
        self.validated_titles_canvas.configure(yscrollcommand=self.validated_titles_scrollbar.set)
        self.validated_titles_canvas.grid(row=0, column=0, sticky="nsew")
        self.validated_titles_scrollbar.grid(row=0, column=1, sticky="ns")

        self.validated_titles_inner = ttk.Frame(self.validated_titles_canvas, padding=(0, 0, 8, 0))
        self.validated_titles_window = self.validated_titles_canvas.create_window((0, 0), window=self.validated_titles_inner, anchor="nw")

        self.validated_titles_inner.bind("<Configure>", self._update_validated_titles_scrollregion)
        self.validated_titles_canvas.bind("<Configure>", self._resize_validated_titles_inner)
        self._refresh_validated_titles_panel()

        ttk.Label(right_panel, text="Textes italiques", font=("Helvetica", 13, "bold")).grid(row=2, column=0, sticky="w", pady=(12, 8))

        italics_container = ttk.Frame(right_panel)
        italics_container.grid(row=3, column=0, sticky="nsew")
        italics_container.columnconfigure(0, weight=1)
        italics_container.rowconfigure(0, weight=1)

        self.validated_italics_canvas = tk.Canvas(italics_container, highlightthickness=0, bg="#ffffff")
        self.validated_italics_scrollbar = ttk.Scrollbar(italics_container, orient="vertical", command=self.validated_italics_canvas.yview)
        self.validated_italics_canvas.configure(yscrollcommand=self.validated_italics_scrollbar.set)
        self.validated_italics_canvas.grid(row=0, column=0, sticky="nsew")
        self.validated_italics_scrollbar.grid(row=0, column=1, sticky="ns")

        self.validated_italics_inner = ttk.Frame(self.validated_italics_canvas, padding=(0, 0, 8, 0))
        self.validated_italics_window = self.validated_italics_canvas.create_window((0, 0), window=self.validated_italics_inner, anchor="nw")

        self.validated_italics_inner.bind("<Configure>", self._update_validated_italics_scrollregion)
        self.validated_italics_canvas.bind("<Configure>", self._resize_validated_italics_inner)
        self._refresh_validated_italics_panel()

        bottom_bar = ttk.Frame(main)
        bottom_bar.grid(row=2, column=0, sticky="ew", pady=(8, 0))
        bottom_bar.rowconfigure(0, weight=0)
        bottom_bar.columnconfigure(3, weight=1)

        legend_frame = tk.Frame(bottom_bar, bd=0, highlightthickness=0, bg=self.cget("bg"))
        legend_frame.grid(row=0, column=0, columnspan=4, sticky="ew", pady=(0, 8))

        legend_items = [
            ("titles", "genre"),
            ("localisation", "localisation"),
            ("bicolumn_text", "bi-colonne"),
            ("italic_text", "italique"),
        ]

        for index, (rect_label, display_label) in enumerate(legend_items):
            item_frame = tk.Frame(legend_frame, bd=0, highlightthickness=0, bg=self.cget("bg"))
            item_frame.pack(side="left", padx=(0 if index == 0 else 18, 0))

            swatch = tk.Canvas(item_frame, width=12, height=12, bd=0, highlightthickness=0, bg=self.cget("bg"))
            swatch.create_oval(2, 2, 10, 10, fill=RECT_COLORS[rect_label], outline=RECT_COLORS[rect_label])
            swatch.pack(side="left")

            ttk.Label(item_frame, text=display_label).pack(side="left", padx=(6, 0))

        ttk.Label(bottom_bar, text="Type:").grid(row=1, column=0, padx=(0, 4))
        self.label_selector = ttk.Combobox(
            bottom_bar,
            textvariable=self.new_rect_label,
            values=["genre", "localisation", "italique", "bi-colonne"],
            state="readonly",
            width=16,
        )
        self.label_selector.grid(row=1, column=1, padx=(0, 12))
        self.label_selector.bind("<<ComboboxSelected>>", self._on_label_selector_change)

        ttk.Button(bottom_bar, text="Ajouter un rectangle", command=self.start_add_rectangle).grid(row=1, column=2, padx=(0, 6))

        self.validate_border_button_text = tk.StringVar(value="Valider bordure page 0")
        ttk.Button(bottom_bar, textvariable=self.validate_border_button_text, command=self.validate_current_page_borders).grid(row=1, column=3, sticky="e", padx=(0, 0))
        ttk.Button(bottom_bar, textvariable=self.export_button_text, command=self.export_validated_data_excel).grid(row=1, column=4, sticky="e", padx=(8, 0))

    def _set_status(self, text):
        return None

    def _page_has_title_rectangle(self, rectangles):
        return any(rect.get("label") == "titles" for rect in rectangles or [])

    def _can_validate_current_page(self):
        current_page_number = self.current_page + 1

        if self._last_validated_border_page_index is None:
            if not self._page_has_title_rectangle(self.current_page_data):
                message = f"Première validation impossible: aucun titre détecté sur la page {current_page_number}."
                self._report_error(message)
                return False
            return True

        expected_page_number = self._last_validated_border_page_index + 2
        if current_page_number != expected_page_number:
            message = f"Page {current_page_number} non attendue. Validez d'abord la page {expected_page_number}."
            self._report_error(message)
            return False

        return True

    def _commit_validated_border_page(self):
        self._register_validated_title(self.current_gallery_sections)
        self._last_validated_border_page_index = self.current_page

    def _prompt_for_pdf(self):
        pdf_path = filedialog.askopenfilename(
            title="Choisir un PDF",
            filetypes=[("PDF files", "*.pdf"), ("All files", "*.*")],
        )
        if pdf_path:
            self.load_pdf(pdf_path)

    def load_pdf(self, pdf_path):
        try:
            self.document = pdfium.PdfDocument(pdf_path)
        except Exception as exc:
            self._report_error(f"Impossible d'ouvrir le PDF: {exc}")
            return

        self.pdf_path = pdf_path
        self.page_count = len(self.document)
        self.current_page = 0
        self.page_annotations = [None] * self.page_count
        self.page_cache = {}
        self.loading_pages = set()
        self._last_validated_border_page_index = None

        self._clear_interaction_state()
        self.current_image = None
        self.current_page_data = []
        self._redraw_canvas()
        self._update_page_label()
        self._request_page_load(0, show_when_ready=True)

    def _request_page_load(self, page_index, show_when_ready=False):
        if not self.pdf_path or not (0 <= page_index < self.page_count):
            return

        if page_index in self.page_cache:
            if show_when_ready:
                self._apply_loaded_page(page_index)
            return

        if page_index in self.loading_pages:
            return

        self.loading_pages.add(page_index)

        worker = threading.Thread(
            target=self._load_page_worker,
            args=(page_index,),
            daemon=True,
        )
        sections = []
        worker.start()

        if show_when_ready and page_index == self.current_page:
            self.current_image = None
            self.current_page_data = []
            self.canvas_placeholder_text = f"Chargement de la page {page_index + 1}"
            self._redraw_canvas()

    def _load_page_worker(self, page_index):
        try:
            document = pdfium.PdfDocument(self.pdf_path)
            image = document[page_index].render(scale=self.render_scale).to_pil().convert("RGB")
            page_result = detect_page_regions(self.pdf_path, page_index=page_index, render_scale=self.render_scale)

            annotations = []
            for label in ("titles", "bicolumn_text", "italic_text"):
                for rect_index, rect in enumerate(page_result.get(label, [])):
                    rect_label = "localisation" if label == "titles" and rect_index == 0 else label
                    annotation = {
                        "x": int(rect["x"]),
                        "y": int(rect["y"]),
                        "w": int(rect["w"]),
                        "h": int(rect["h"]),
                        "label": rect_label,
                    }

                    if label == "bicolumn_text" and "center_x" in rect:
                        annotation["center_x"] = int(rect["center_x"])

                    annotations.append(annotation)

            self.load_queue.put(("page_loaded", page_index, image, image.size, annotations))
        except Exception as exc:
            self.load_queue.put(("page_error", page_index, str(exc)))

    def _poll_load_queue(self):
        try:
            while True:
                message = self.load_queue.get_nowait()
                event = message[0]

                if event == "page_loaded":
                    _, page_index, image, source_size, annotations = message
                    self.page_cache[page_index] = {
                        "image": image,
                        "source_size": source_size,
                        "annotations": annotations,
                    }
                    self.page_annotations[page_index] = annotations
                    self.loading_pages.discard(page_index)

                    if page_index == self.current_page:
                        self._apply_loaded_page(page_index)
                elif event == "page_error":
                    _, page_index, error_message = message
                    self.loading_pages.discard(page_index)
                    if page_index == self.current_page:
                        self.current_image = None
                        self.current_page_data = []
                        self.canvas_placeholder_text = "Aucune page à afficher"
                        self._redraw_canvas()
                        self._set_status(f"Erreur page {page_index + 1}")
                    self._report_error(f"Impossible de charger la page {page_index + 1}: {error_message}")
        except queue.Empty:
            pass

        self._poll_after_id = self.after(100, self._poll_load_queue)

    def _apply_loaded_page(self, page_index):
        cached = self.page_cache.get(page_index)
        if not cached:
            return

        self.current_page = page_index
        self.page_jump_var.set(str(page_index + 1))
        self.source_image_size = cached.get("source_size", cached["image"].size)
        self.current_image = cached["image"]
        self.current_page_data = self._ensure_rect_uids(cached["annotations"])
        cached["annotations"] = self.current_page_data
        self.page_annotations[page_index] = self.current_page_data
        self.canvas_placeholder_text = "Aucune page à afficher"
        self._clear_interaction_state()
        self._redraw_canvas()
        self._update_page_label()

    def _new_rect_uid(self):
        uid = self._next_rect_uid
        self._next_rect_uid += 1
        return uid

    def _ensure_rect_uids(self, rectangles):
        normalized = []
        for rect in rectangles:
            item = rect.copy()
            if not item.get("uid"):
                item["uid"] = self._new_rect_uid()
            else:
                self._next_rect_uid = max(self._next_rect_uid, int(item["uid"]) + 1)
            normalized.append(item)
        return normalized

    def _load_page_annotations(self):
        # Conservé pour compatibilité, mais le chargement réel est désormais asynchrone.
        return None

    def _show_page(self, page_index):
        if not self.document or not (0 <= page_index < self.page_count):
            return

        self._save_current_page_state()
        self.current_page = page_index
        self.page_jump_var.set(str(page_index + 1))
        self._update_page_label()

        if page_index in self.page_cache:
            self._apply_loaded_page(page_index)
        else:
            self.current_image = None
            self.current_page_data = []
            self._clear_interaction_state()
            self.canvas_placeholder_text = f"Chargement de la page {page_index + 1}..."
            self._redraw_canvas()
            self._request_page_load(page_index, show_when_ready=True)

    def _save_current_page_state(self):
        if self.page_annotations and 0 <= self.current_page < len(self.page_annotations):
            self.page_annotations[self.current_page] = [rect.copy() for rect in self.current_page_data]
            if self.current_page in self.page_cache:
                self.page_cache[self.current_page]["annotations"] = self.page_annotations[self.current_page]

    def _update_validated_titles_scrollregion(self, event=None):
        if hasattr(self, "validated_titles_canvas") and self.validated_titles_canvas.winfo_exists():
            self.validated_titles_canvas.configure(scrollregion=self.validated_titles_canvas.bbox("all"))

    def _resize_validated_titles_inner(self, event):
        if hasattr(self, "validated_titles_canvas") and self.validated_titles_canvas.winfo_exists():
            self.validated_titles_canvas.itemconfigure(self.validated_titles_window, width=event.width)

    def _clear_validated_titles_panel(self):
        self._clear_children(self.validated_titles_inner)

    def _refresh_validated_titles_panel(self):
        if not hasattr(self, "validated_titles_inner"):
            return

        self._clear_validated_titles_panel()

        if not self.validated_titles:
            ttk.Label(self.validated_titles_inner, text="Aucun genre validé", foreground="#6b7280").pack(anchor="w", pady=(0, 8))
            self._update_validated_titles_scrollregion()
            return

        for title_group in self.validated_titles:
            button = ttk.Button(
                self.validated_titles_inner,
                text=self._format_validated_title_button(title_group),
                command=lambda group=title_group: self._open_validated_title_detail(group),
            )
            button.pack(fill="x", pady=(0, 6))

        self._update_validated_titles_scrollregion()

    def _update_validated_italics_scrollregion(self, event=None):
        if hasattr(self, "validated_italics_canvas") and self.validated_italics_canvas.winfo_exists():
            self.validated_italics_canvas.configure(scrollregion=self.validated_italics_canvas.bbox("all"))

    def _resize_validated_italics_inner(self, event):
        if hasattr(self, "validated_italics_canvas") and self.validated_italics_canvas.winfo_exists():
            self.validated_italics_canvas.itemconfigure(self.validated_italics_window, width=event.width)

    def _clear_validated_italics_panel(self):
        if hasattr(self, "validated_italics_inner"):
            self._clear_children(self.validated_italics_inner)

    def _refresh_validated_italics_panel(self):
        if not hasattr(self, "validated_italics_inner"):
            return

        self._clear_validated_italics_panel()

        if not self.validated_italics:
            ttk.Label(self.validated_italics_inner, text="Aucun texte italique validé", foreground="#6b7280").pack(anchor="w", pady=(0, 8))
            self._update_validated_italics_scrollregion()
            return

        self._validated_italics_preview_refs = []

        for italic_item in self.validated_italics:
            card = ttk.LabelFrame(self.validated_italics_inner, text=italic_item.get("title") or "Texte italique", padding=8)
            card.pack(fill="x", pady=(0, 8))
            card.columnconfigure(1, weight=1)

            with Image.open(italic_item["image_path"]) as italic_image:
                preview = italic_image.convert("RGB")
                preview.thumbnail((220, 140), Image.Resampling.LANCZOS)
                photo = ImageTk.PhotoImage(preview)

            self._validated_italics_preview_refs.append(photo)

            ttk.Label(card, image=photo).grid(row=0, column=0, rowspan=3, sticky="nw", padx=(0, 10))
            ttk.Label(card, text=italic_item.get("title", "Texte italique"), font=("Helvetica", 11, "bold")).grid(row=0, column=1, sticky="w")

            localisation = italic_item.get("localisation", "")
            ttk.Label(card, text=f"Localisation : {localisation or '—'}", foreground="#4b5563").grid(row=1, column=1, sticky="w", pady=(4, 0))
            ttk.Button(card, text="Ouvrir", command=lambda item=italic_item: self._open_validated_italic_detail(item)).grid(row=2, column=1, sticky="w", pady=(8, 0))

        self._update_validated_italics_scrollregion()

    def _build_italic_section(self, parent, item):
        card = ttk.LabelFrame(parent, text="texte italique", padding=10)
        card.pack(fill="x", pady=(0, 12))
        card.columnconfigure(0, weight=1)

        ttk.Label(card, text=item["name"], font=("Helvetica", 11, "bold")).grid(row=0, column=0, sticky="w", pady=(0, 8))

        preview = None
        with Image.open(item["path"]) as italic_image:
            preview = italic_image.convert("RGB")
            preview.thumbnail((900, 260), Image.Resampling.LANCZOS)

        photo = ImageTk.PhotoImage(preview)
        card._photo_ref = photo
        ttk.Label(card, image=photo).grid(row=1, column=0, sticky="w")

        section = {
            "kind": "italic",
            "title": "texte italique",
            "image_path": item["path"],
            "source_name": item["name"],
        }
        return section

    def _open_validated_italic_detail(self, italic_item):
        detail_window = tk.Toplevel(self)
        detail_window.title(italic_item.get("title", "Texte italique"))
        detail_window.geometry("1000x700")
        detail_window.columnconfigure(0, weight=1)
        detail_window.rowconfigure(1, weight=1)

        container = ttk.Frame(detail_window, padding=12)
        container.grid(row=0, column=0, sticky="nsew")
        container.columnconfigure(0, weight=1)

        ttk.Label(container, text=italic_item.get("title", "Texte italique"), font=("Helvetica", 15, "bold")).grid(row=0, column=0, sticky="w")
        ttk.Label(container, text=f"Localisation : {italic_item.get('localisation', '')}", font=("Helvetica", 11)).grid(row=1, column=0, sticky="w", pady=(6, 12))

        image_frame = ttk.Frame(detail_window, padding=12)
        image_frame.grid(row=1, column=0, sticky="nsew")
        image_frame.columnconfigure(0, weight=1)
        image_frame.rowconfigure(0, weight=1)

        with Image.open(italic_item["image_path"]) as italic_image:
            preview = italic_image.convert("RGB")
            preview.thumbnail((950, 500), Image.Resampling.LANCZOS)
            photo = ImageTk.PhotoImage(preview)

        detail_window._photo_ref = photo
        ttk.Label(image_frame, image=photo).grid(row=0, column=0, sticky="w")

    def _build_export_title_lines(self, node, lines, level=0):
        indent = "    " * level
        node_type = node.get("type", "taxon")
        node_text = (node.get("text", "") or "").strip()

        if node_type == "title":
            lines.append(f"{indent}titre : {node_text}")
        elif node_type == "taxon":
            taxon_text, separator, reference_text = node_text.partition("%")
            lines.append(f"{indent}taxon : {taxon_text.strip()}")
            if separator:
                lines.append(f"{indent}    reference : {reference_text.strip()}")
        else:
            lines.append(f"{indent}{node_type} : {node_text}")

        for child in node.get("children", []):
            self._build_export_title_lines(child, lines, level + 1)

    def _split_taxon_reference_text(self, raw_text):
        taxon_text, separator, reference_text = (raw_text or "").partition("%")
        return taxon_text.strip(), reference_text.strip() if separator else ""

    def _export_validated_titles_rows(self):
        rows = []

        def walk_nodes(nodes, title_text, current_espece=None, current_variete=None):
            for node in nodes or []:
                node_type = node.get("type", "")
                node_text = (node.get("text", "") or "").strip()

                if node_type in {"espèce", "espece"}:
                    current_espece = node_text
                elif node_type in {"variété", "variete"}:
                    current_variete = node_text
                elif node_type == "taxon":
                    taxon_text, reference_text = self._split_taxon_reference_text(node_text)
                    rows.append(
                        {
                            "taxon": taxon_text,
                            "reference": reference_text,
                            "espèce": current_espece or "",
                            "variété": current_variete or "",
                            "genre": title_text,
                        }
                    )

                current_espece, current_variete = walk_nodes(
                    node.get("children", []),
                    title_text,
                    current_espece,
                    current_variete,
                )

            return current_espece, current_variete

        for title_group in self.validated_titles:
            title_text = title_group.get("title", "")
            walk_nodes(title_group.get("hierarchy", []), title_text)

        return rows

    def _get_reference_dictionary_for_export(self):
        if self.reference_dictionary is not None:
            return self.reference_dictionary

        reference_path = filedialog.askopenfilename(
            title="Choisir le fichier des références",
            filetypes=[("Fichier Excel", "*.xlsx"), ("Tous les fichiers", "*.*")],
        )
        if not reference_path:
            return None

        try:
            self.reference_dictionary = reference_fuzzy.load_reference_from_xlsx(reference_path)
            self.reference_dictionary_path = reference_path
            self.export_button_text.set("Exporter Excel")
        except Exception as exc:
            self._report_error(f"Impossible de charger le fichier des références: {exc}")
            return None

        return self.reference_dictionary

    def export_validated_data_excel(self):
        if not self.validated_titles and not self.validated_italics:
            self._report_error("Aucune donnée validée à exporter.")
            return

        reference_dictionary = self._get_reference_dictionary_for_export()
        if reference_dictionary is None:
            return

        default_name = "export_validations.xlsx"
        if self.pdf_path:
            default_name = f"{Path(self.pdf_path).stem}_validations.xlsx"

        export_path = filedialog.asksaveasfilename(
            title="Enregistrer l'export Excel",
            defaultextension=".xlsx",
            initialfile=default_name,
            filetypes=[("Fichier Excel", "*.xlsx")],
        )
        if not export_path:
            return

        workbook = Workbook()
        titles_sheet = workbook.active
        titles_sheet.title = "Genres validés"
        italics_sheet = workbook.create_sheet("Textes italiques")

        title_headers = ["taxon", "reference", "espèce", "variété", "genre"]
        titles_sheet.append(title_headers)
        for row in self._export_validated_titles_rows():
            fuzzy_reference = reference_fuzzy.find_fuzzy_references(row["reference"], reference_dictionary=reference_dictionary)
            titles_sheet.append([row["taxon"], fuzzy_reference, row["espèce"], row["variété"], row["genre"]])

        italic_headers = ["id", "image", "titre", "localisation"]
        italics_sheet.append(italic_headers)
        exported_italic_images = []
        max_italic_image_width = 0
        for index, italic_item in enumerate(self.validated_italics, start=1):
            row_number = index + 1
            italics_sheet.append(
                [
                    index,
                    "",
                    italic_item.get("title", ""),
                    italic_item.get("localisation", ""),
                ]
            )

            image_path = italic_item.get("image_path") or italic_item.get("path")
            if image_path:
                with Image.open(image_path) as source_image:
                    source_width, source_height = source_image.size
                    scaled_width = max(1, int(source_width / 10))
                    scaled_height = max(1, int(source_height / 10))
                    max_italic_image_width = max(max_italic_image_width, scaled_width)

                excel_image = OpenpyxlImage(str(image_path))
                excel_image.width = scaled_width
                excel_image.height = scaled_height
                excel_image.anchor = f"B{row_number}"
                italics_sheet.add_image(excel_image)
                exported_italic_images.append(excel_image)

                italics_sheet.row_dimensions[row_number].height = max(18, scaled_height * 0.75)

        titles_sheet.column_dimensions["A"].width = 30
        titles_sheet.column_dimensions["B"].width = 40
        titles_sheet.column_dimensions["C"].width = 60
        titles_sheet.column_dimensions["D"].width = 40

        italics_sheet.column_dimensions["A"].width = 8
        italics_sheet.column_dimensions["B"].width = max(12, max_italic_image_width / 7.0)
        italics_sheet.column_dimensions["C"].width = 32
        italics_sheet.column_dimensions["D"].width = 45

        for sheet in (titles_sheet, italics_sheet):
            for column_cells in sheet.columns:
                max_length = 0
                column_letter = column_cells[0].column_letter
                if sheet.title == "Textes italiques" and column_letter == "B":
                    continue
                for cell in column_cells:
                    value = "" if cell.value is None else str(cell.value)
                    max_length = max(max_length, max((len(line) for line in value.splitlines()), default=0))
                sheet.column_dimensions[column_letter].width = min(max_length + 2, 80)
                for cell in column_cells:
                    cell.alignment = cell.alignment.copy(wrap_text=True, vertical="top")

        workbook.save(export_path)
        messagebox.showinfo("Export terminé", f"Fichier Excel créé :\n{export_path}")

    def _format_validated_title_button(self, title_group):
        return title_group.get('title', 'Genre')

    def _extract_title_group_name(self, exported_items):
        for item in exported_items:
            normalized_name = item["name"].lower().replace("-", "_")
            if normalized_name.startswith("genre") or normalized_name.startswith("titre"):
                _, fuzzy_title_text = self._ocr_title_text(item["path"])
                cleaned_title = " ".join((fuzzy_title_text or item["name"]).split())
                if cleaned_title:
                    return cleaned_title
        return None

    def _build_title_hierarchy(self, exported_items):
        hierarchy = []
        pending_page_localisations = []
        pending_page_italics = []
        pending_title_italics = []
        current_title = None

        for item in exported_items:
            normalized_name = item["name"].lower().replace("-", "_")

            if is_italic_name(normalized_name):
                pending_page_italics.append({"path": item["path"]})
                if current_title is not None:
                    current_title.setdefault("italics", []).append({"path": item["path"]})
                else:
                    pending_title_italics.append({"path": item["path"]})
                continue

            if "localisation" in normalized_name:
                pending_page_localisations.append(self._ocr_image_text(item["path"]))
                continue

            if normalized_name.startswith("genre") or normalized_name.startswith("titre"):
                raw_title_text, fuzzy_title_text = self._ocr_title_text(item["path"])
                current_title = {
                    "type": "title",
                    "label": "titre",
                    "text": fuzzy_title_text or raw_title_text or item["name"],
                    "italics": pending_title_italics[:],
                    "children": [],
                }
                hierarchy.append(current_title)
                pending_title_italics = []
                continue

            if ("bicolonne" in normalized_name or "bi_colonne" in normalized_name) and hierarchy:
                current_title = hierarchy[-1]
                for block in self._ocr_bicolumn_blocks(item["path"]):
                    current_title["children"].append(
                        {
                            "type": self._ocr_type_display(block.get("type", "taxon")),
                            "text": block.get("text", ""),
                        }
                    )
                continue

            if hierarchy:
                current_title = hierarchy[-1]
                current_title["children"].append(
                    {
                        "type": "taxon",
                        "text": self._ocr_image_text(item["path"]),
                    }
                )

        return {"titles": hierarchy, "localisations": pending_page_localisations, "italics": pending_page_italics}

    def _register_validated_title(self, gallery_sections):
        current_group = self.validated_titles[-1] if self.validated_titles else None
        current_title_entry = current_group["hierarchy"][-1] if current_group and current_group.get("hierarchy") else None
        page_localizations = []  # Initialize the list for localizations
        italic_items = []  # Initialize the list for italic items

        def ensure_current_group(title_text):
            nonlocal current_group, current_title_entry

            title_text = " ".join((title_text or "").split())
            if not title_text:
                title_text = f"Genre validé {self._validated_title_counter}"

            current_group = {
                "id": self._validated_title_counter,
                "title": title_text,
                "pages": [],
                "hierarchy": [
                    {
                        "type": "title",
                        "label": "titre",
                        "text": title_text,
                        "children": [],
                    }
                ],
            }
            self._validated_title_counter += 1
            self.validated_titles.append(current_group)
            current_title_entry = current_group["hierarchy"][0]
            return current_group, current_title_entry

        def ensure_page_entry(group):
            page_key = self.current_page
            pages = group.setdefault("pages", [])
            if not pages or pages[-1].get("page_index") != page_key:
                pages.append(
                    {
                        "page_index": self.current_page,
                        "page_number": self.current_page + 1,
                        "items": gallery_sections,
                    }
                )
            return pages[-1]

        def append_children(children):
            nonlocal current_group, current_title_entry
            if not children:
                return

            if current_group is None:
                if self.validated_titles:
                    current_group = self.validated_titles[-1]
                    current_title_entry = current_group["hierarchy"][-1] if current_group.get("hierarchy") else None
                else:
                    return

            if current_title_entry is None:
                current_group, current_title_entry = ensure_current_group(f"Genre validé {self._validated_title_counter}")

            current_title_entry.setdefault("children", []).extend(children)
            ensure_page_entry(current_group)

        for section in gallery_sections or []:
            if section is None:
                continue

            kind = section.get("kind")
            if kind == "italic":
                italic_record = {
                    "image_path": section.get("image_path"),
                    "source_name": section.get("source_name", "texte italique"),
                    "title": current_group.get("title") if current_group else None,
                }
                if current_group is not None:
                    current_group.setdefault("italics", []).append(italic_record)
                italic_items.append(italic_record)
                continue

            if kind == "simple":
                self._sync_ocr_section_rows(section)
                rows = [row for row in section.get("rows", []) if row.get("text", "").strip()]
                section_title = section.get("title")

                if section_title == "localisation":
                    page_localizations.extend(row.get("text", "") for row in rows)
                    continue

                if section_title in {"titre", "genre"} and rows:
                    title_text = rows[0].get("text", "")
                    current_group, current_title_entry = ensure_current_group(title_text)
                    ensure_page_entry(current_group)
                    continue

                children = [
                    {
                        "type": self._ocr_type_display(section.get("fixed_type", "taxon")),
                        "text": "\n".join(row.get("text", "") for row in rows),
                    }
                ] if rows else []
                append_children(children)
                continue

            if kind == "bicolumn":
                children = []
                for side in section.get("sides", []):
                    if side is None:
                        continue
                    self._sync_ocr_section_rows(side)
                    for row in side.get("rows", []):
                        row_text = row.get("text", "").strip()
                        if row_text:
                            children.append({"type": self._ocr_type_display(row.get("type", "taxon")), "text": row_text})

                append_children(children)

        page_localization_text = "\n".join(text for text in page_localizations if text).strip()

        for italic_item in italic_items:
            italic_item["localisation"] = page_localization_text
            if italic_item.get("title") is None and current_group is not None:
                italic_item["title"] = current_group.get("title")
            self.validated_italics.append(italic_item)

        self._refresh_validated_titles_panel()
        self._refresh_validated_italics_panel()

    def _open_validated_title_detail(self, title_group):
        detail_window = tk.Toplevel(self)
        detail_window.title(title_group.get("title", "Genre validé"))
        detail_window.geometry("1100x800")
        detail_window.columnconfigure(0, weight=1)
        detail_window.rowconfigure(0, weight=1)

        container = ttk.Frame(detail_window, padding=12)
        container.grid(row=0, column=0, sticky="nsew")
        container.columnconfigure(0, weight=1)
        container.rowconfigure(1, weight=1)

        ttk.Label(container, text=title_group.get("title", "Genre validé"), font=("Helvetica", 15, "bold")).grid(row=0, column=0, sticky="w", pady=(0, 10))

        text = tk.Text(container, wrap="word")
        text.grid(row=1, column=0, sticky="nsew")
        text.tag_configure("level0", lmargin1=0, lmargin2=0)
        text.tag_configure("level1", lmargin1=28, lmargin2=28)
        text.tag_configure("level2", lmargin1=56, lmargin2=56)
        text.tag_configure("level3", lmargin1=84, lmargin2=84)
        text.tag_configure("level4", lmargin1=112, lmargin2=112)

        scrollbar = ttk.Scrollbar(container, orient="vertical", command=text.yview)
        scrollbar.grid(row=1, column=1, sticky="ns")
        text.configure(yscrollcommand=scrollbar.set)

        def append(line, level="level0"):
            text.insert("end", f"{line}\n", level)

        def append_taxon_with_reference(raw_text, level="level3"):
            taxon_text, separator, reference_text = (raw_text or "").partition("%")
            append(f"taxon : {taxon_text.strip()}", level=level)
            if separator:
                append(f"reference : {reference_text.strip()}", level="level4")

        preview_refs = []

        for title in title_group.get("hierarchy", []):
            append(f"genre : {title.get('text', '')}", level="level0")

            for child in title.get("children", []):
                child_type = child.get("type", "taxon")
                child_indent = {
                    "espèce": "level1",
                    "espece": "level1",
                    "variété": "level2",
                    "variete": "level2",
                    "taxon": "level3",
                }.get(child_type, "level3")
                if child_type == "taxon":
                    append_taxon_with_reference(child.get("text", ""), level=child_indent)
                else:
                    append(f"{child_type} : {child.get('text', '')}", level=child_indent)

            append("", level="level0")

        italics = title_group.get("italics", [])
        if italics:
            italics_frame = ttk.LabelFrame(container, text="Textes italiques", padding=10)
            italics_frame.grid(row=2, column=0, columnspan=2, sticky="ew", pady=(10, 0))
            italics_frame.columnconfigure(1, weight=1)

            for italic_index, italic_item in enumerate(italics, start=1):
                row_frame = ttk.Frame(italics_frame)
                row_frame.grid(row=italic_index - 1, column=0, sticky="ew", pady=(0, 8))
                row_frame.columnconfigure(1, weight=1)

                image_path = italic_item.get("image_path") or italic_item.get("path")
                if image_path:
                    with Image.open(image_path) as italic_image:
                        preview = italic_image.convert("RGB")
                        preview.thumbnail((220, 140), Image.Resampling.LANCZOS)
                        photo = ImageTk.PhotoImage(preview)

                    preview_refs.append(photo)
                    ttk.Label(row_frame, image=photo).grid(row=0, column=0, rowspan=2, sticky="nw", padx=(0, 10))

                ttk.Label(row_frame, text=f"italique : {italic_item.get('source_name', 'texte italique')}").grid(row=0, column=1, sticky="w")
                ttk.Label(row_frame, text=f"Localisation : {italic_item.get('localisation', '—')}", foreground="#4b5563").grid(row=1, column=1, sticky="w")

        detail_window._preview_refs = preview_refs

        text.configure(state="disabled")

    def _update_page_label(self):
        self.page_label.config(text=f"Page {self.current_page + 1 if self.page_count else 0} / {self.page_count}")
        self.validate_border_button_text.set(f"Valider bordure page {self.current_page + 1 if self.page_count else 0}")

    def _make_export_display_name(self, rect_label, count, variant=None):
        display_label = label_to_display(rect_label)

        if rect_label == "localisation":
            return display_label if count == 1 else f"{display_label} {count}"

        if rect_label == "bicolumn_text" and variant:
            return f"{display_label} {count} {variant}"

        return f"{display_label} {count}"

    def _export_current_page_rectangles(self):
        if self.current_image is None:
            self._report_error("Aucune page affichée à enregistrer.")
            return []

        rectangles = self.current_page_data or []
        if not rectangles:
            self._report_error("Aucun rectangle à enregistrer sur cette page.")
            return []

        if self.pdf_path:
            debug_dir = Path(self.pdf_path).with_name(f"{Path(self.pdf_path).stem}_page_{self.current_page + 1}_rectangles")
        else:
            debug_dir = Path.cwd() / f"page_{self.current_page + 1}_rectangles"

        debug_dir.mkdir(parents=True, exist_ok=True)

        saved_items = []
        image_width, image_height = self.current_image.size
        label_counts = {"titles": 0, "localisation": 0, "italic_text": 0, "bicolumn_text": 0}

        ordered_rectangles = sorted(rectangles, key=lambda rect: (rect["y"], rect["x"]))

        for index, rect in enumerate(ordered_rectangles, start=1):
            x1 = int(clamp(rect["x"], 0, image_width))
            y1 = int(clamp(rect["y"], 0, image_height))
            x2 = int(clamp(rect["x"] + rect["w"], 0, image_width))
            y2 = int(clamp(rect["y"] + rect["h"], 0, image_height))

            if x2 <= x1 or y2 <= y1:
                continue

            rect_uid = rect.get("uid", index)
            rect_label = rect.get("label", "manual")

            if rect_label == "bicolumn_text":
                label_counts[rect_label] += 1
                rect_count = label_counts[rect_label]
                separator_x = int(clamp(rect.get("center_x", x1 + (x2 - x1) // 2), x1, x2))

                if separator_x > x1:
                    left_name = self._make_export_display_name(rect_label, rect_count, "gauche")
                    left_path = debug_dir / f"{left_name.lower().replace(' ', '_')}.png"
                    self.current_image.crop((x1, y1, separator_x, y2)).save(left_path)
                    saved_items.append({"path": left_path, "name": left_name})

                if x2 > separator_x:
                    right_name = self._make_export_display_name(rect_label, rect_count, "droite")
                    right_path = debug_dir / f"{right_name.lower().replace(' ', '_')}.png"
                    self.current_image.crop((separator_x, y1, x2, y2)).save(right_path)
                    saved_items.append({"path": right_path, "name": right_name})
                continue

            label_counts[rect_label] = label_counts.get(rect_label, 0) + 1
            rect_count = label_counts[rect_label]
            display_name = self._make_export_display_name(rect_label, rect_count)
            export_path = debug_dir / f"{display_name.lower().replace(' ', '_')}.png"
            self.current_image.crop((x1, y1, x2, y2)).save(export_path)
            saved_items.append({"path": export_path, "name": display_name})

        if not saved_items:
            self._report_error("Aucun rectangle valide n'a pu être enregistré.")
            return []

        return saved_items

    def _normalize_ocr_text(self, text):
        cleaned = delete_bars(text or "")
        cleaned = "\n".join(line.rstrip() for line in cleaned.splitlines())
        return cleaned.strip()

    def _strip_first_word(self, text):
        lines = (text or "").splitlines()
        if not lines:
            return ""

        first_line = lines[0].lstrip()
        parts = first_line.split(maxsplit=1)
        if len(parts) < 2:
            cleaned_lines = lines[1:]
        else:
            cleaned_lines = [parts[1], *lines[1:]]

        return "\n".join(cleaned_lines).strip()

    def _ocr_image_text(self, image_path):
        with Image.open(image_path) as image:
            try:
                raw_text = pytesseract.image_to_string(image, lang="lat")
            except Exception:
                raw_text = pytesseract.image_to_string(image)
        return self._normalize_ocr_text(raw_text)

    def _ocr_title_with_tesseract_variants(self, image):
        variants = [
            image,
            ImageOps.autocontrast(image.convert("L")),
            ImageOps.autocontrast(image.convert("L")).resize(
                (max(1, image.width * 2), max(1, image.height * 2)),
                Image.Resampling.LANCZOS,
            ),
            ImageOps.autocontrast(image.convert("L")).resize(
                (max(1, image.width * 3), max(1, image.height * 3)),
                Image.Resampling.LANCZOS,
            ),
        ]

        configs = [
            "--oem 3 --psm 6",
            "--oem 3 --psm 7",
            "--oem 3 --psm 11",
        ]

        for variant in variants:
            for config in configs:
                try:
                    raw_text = pytesseract.image_to_string(variant, lang="lat", config=config)
                except Exception:
                    try:
                        raw_text = pytesseract.image_to_string(variant, config=config)
                    except Exception:
                        continue

                cleaned_text = self._normalize_ocr_text(raw_text)
                if cleaned_text:
                    return cleaned_text

        return ""

    def _ocr_title_text(self, image_path):
        with Image.open(image_path) as image:
            raw_text = self._ocr_title_with_tesseract_variants(image)

        if not raw_text:
            raw_text = self._ocr_image_text(image_path)

        fuzzy_text = remove_fuzzy_phrase(
            raw_text,
            ["EJUSQUE SPECIES", "ET EJUS SPECIES"],
        )
        return raw_text, fuzzy_text

    def _ocr_bicolumn_blocks(self, image_path):
        with Image.open(image_path) as image:
            analysis = extract_line_indentation(image)

        if not analysis or not analysis.get("blocks"):
            return []

        blocks = []
        for block in analysis["blocks"]:
            if block.get("type") == "commentaire":
                continue

            text = self._normalize_ocr_text(block.get("text", ""))
            if block.get("type") == "variete":
                text = self._strip_first_word(text)
            elif block.get("type", "taxon") == "taxon":
                text = replace_first_taxon_punctuation(text)
            if not text:
                continue

            blocks.append(
                {
                    "type": block.get("type", "taxon"),
                    "text": text,
                }
            )

        return blocks

    def _create_text_widget(self, parent, text="", height=3):
        text_widget = tk.Text(parent, height=height, wrap="word", relief="solid", borderwidth=1)
        text_widget.insert("1.0", text)
        text_widget.tag_configure("percent", foreground="#d62728")

        def refresh_percent_highlight(event=None, widget=text_widget):
            widget.after_idle(lambda w=widget: self._highlight_percent_in_text_widget(w))

        text_widget.bind("<KeyRelease>", refresh_percent_highlight)
        text_widget.bind("<ButtonRelease-1>", refresh_percent_highlight)
        text_widget.bind("<<Paste>>", refresh_percent_highlight)
        self._highlight_percent_in_text_widget(text_widget)
        return text_widget

    def _highlight_percent_in_text_widget(self, text_widget):
        if text_widget is None or not text_widget.winfo_exists():
            return

        text_widget.tag_remove("percent", "1.0", "end")
        start_index = "1.0"

        while True:
            percent_index = text_widget.search("%", start_index, stopindex="end")
            if not percent_index:
                break

            text_widget.tag_add("percent", percent_index, f"{percent_index}+1c")
            start_index = f"{percent_index}+1c"

    def _ocr_type_display(self, ocr_type):
        return {
            "espece": "espèce",
            "variete": "variété",
            "taxon": "taxon",
            "titre": "genre",
            "localisation": "localisation",
            "entry": "taxon",
        }.get(ocr_type, ocr_type)

    def _ocr_type_with_bullet(self, ocr_type):
        return f"• {self._ocr_type_display(ocr_type)}"

    def _ocr_type_internal(self, display_type):
        return {
            "espèce": "espece",
            "variété": "variete",
            "taxon": "taxon",
            "genre": "titre",
            "localisation": "localisation",
        }.get(display_type, display_type)

    def _ocr_row_indent(self, row_type):
        return {
            "espece": 0,
            "espèce": 0,
            "variete": 28,
            "variété": 28,
            "taxon": 56,
            "titre": 0,
            "localisation": 0,
            "entry": 56,
        }.get(row_type, 0)

    def _clear_children(self, widget):
        for child in widget.winfo_children():
            child.destroy()

    def _sync_ocr_section_rows(self, section):
        for row in section["rows"]:
            text_widget = row.get("text_widget")
            if text_widget is not None and text_widget.winfo_exists():
                row["text"] = text_widget.get("1.0", "end-1c").strip()

            type_var = row.get("type_var")
            if type_var is not None:
                row["type"] = self._ocr_type_internal(type_var.get())

    def _insert_ocr_row(self, section, index, row_type=None, text=""):
        self._sync_ocr_section_rows(section)

        default_type = row_type or section.get("fixed_type") or "taxon"
        if section.get("kind") == "bicolumn":
            default_type = self._ocr_type_internal(self._ocr_type_display(default_type))

        section["rows"].insert(
            max(0, min(index, len(section["rows"]))),
            {
                "type": default_type,
                "text": text,
            },
        )
        self._render_ocr_section(section)

    def _delete_ocr_row(self, section, index):
        self._sync_ocr_section_rows(section)
        if 0 <= index < len(section["rows"]):
            row = section["rows"][index]
            row_type = self._ocr_type_display(row.get("type", "taxon"))
            row_text = row.get("text", "").strip()
            confirmation_text = f"Êtes-vous sûr de vouloir supprimer {row_type} :\n\n{row_text or '[contenu vide]'}"

            if not messagebox.askyesno("Confirmation de suppression", confirmation_text):
                return

            section["rows"].pop(index)
        self._render_ocr_section(section)

    def _should_show_ocr_add_button(self, section, index):
        rows = section.get("rows", [])
        return 0 <= index < len(rows)

    def _merge_consecutive_taxon_rows(self, rows):
        merged_rows = []

        for row in rows or []:
            row_type = row.get("type", "taxon")
            row_text = row.get("text", "")

            if (
                merged_rows
                and row_type in {"espece", "variete"}
                and merged_rows[-1].get("type") == row_type
            ):
                previous_text = merged_rows[-1].get("text", "")
                merged_rows[-1]["text"] = "\n".join(text for text in (previous_text, row_text) if text).strip()
                continue

            merged_rows.append(row.copy())

        return merged_rows

    def _create_action_button(self, parent, text, bg, active_bg, command):
        container = tk.Frame(parent, bg=bg, width=48, height=28, bd=0, highlightthickness=0)
        container.pack_propagate(False)

        button = tk.Label(
            container,
            text=text,
            fg="white",
            bg=bg,
            bd=0,
            relief="flat",
            cursor="hand2",
            font=("Helvetica", 16, "bold"),
        )
        button.pack(fill="both", expand=True)

        def set_active(state_active):
            color = active_bg if state_active else bg
            container.configure(bg=color)
            button.configure(bg=color)

        container.bind("<Button-1>", lambda event: command())
        container.bind("<Enter>", lambda event: set_active(True))
        container.bind("<Leave>", lambda event: set_active(False))
        button.bind("<Button-1>", lambda event: command())
        button.bind("<Enter>", lambda event: set_active(True))
        button.bind("<Leave>", lambda event: set_active(False))
        return container

    def _render_ocr_section(self, section):
        self._clear_children(section["rows_frame"])

        rows = section["rows"]
        if section.get("kind") == "bicolumn":
            top_insert = ttk.Frame(section["rows_frame"])
            top_insert.pack(fill="x", pady=(0, 8), padx=(8, 8))
            ttk.Label(top_insert, text="", width=14, anchor="w").pack(side="left", padx=(0, 6))
            self._create_action_button(
                top_insert,
                text="+",
                bg="#2ca02c",
                active_bg="#228b22",
                command=lambda s=section: self._insert_ocr_row(s, 0),
            ).pack(side="left")

        if not rows:
            ttk.Label(section["rows_frame"], text="Aucune zone", foreground="#6b7280").pack(anchor="w", pady=(0, 8))

        for index, row in enumerate(rows):
            row_frame = ttk.Frame(section["rows_frame"])
            row_frame.pack(fill="x", pady=(0, 8), padx=(8, 8))

            indent = self._ocr_row_indent(row.get("type")) if section.get("kind") == "bicolumn" else 0
            row_inner = ttk.Frame(row_frame)
            row_inner.pack(fill="x", padx=(indent, 0))
            row_inner.columnconfigure(0, weight=0)
            row_inner.columnconfigure(1, weight=1)
            row_inner.columnconfigure(2, weight=0)

            header = ttk.Frame(row_inner)
            header.grid(row=0, column=0, sticky="w")
            row_inner.columnconfigure(0, weight=0)

            if section.get("kind") == "bicolumn":
                type_var = tk.StringVar(value=self._ocr_type_display(row.get("type", "taxon")))
                type_selector = ttk.Combobox(
                    header,
                    textvariable=type_var,
                    values=["espèce", "variété", "taxon"],
                    state="readonly",
                    width=12,
                )
                ttk.Label(header, text="•", width=2, anchor="w").pack(side="left", padx=(0, 2))
                type_selector.pack(side="left", padx=(0, 6))
                row["type_var"] = type_var
                type_selector.bind(
                    "<<ComboboxSelected>>",
                    lambda event, s=section: (self._sync_ocr_section_rows(s), self._render_ocr_section(s)),
                )
            else:
                ttk.Label(header, text=self._ocr_type_with_bullet(section.get("fixed_type_display", section["title"])), width=14, anchor="w").pack(side="left", padx=(0, 6))
                row["type_var"] = None

            text_widget = self._create_text_widget(row_inner, row.get("text", ""), height=3 if section.get("kind") == "bicolumn" else 4)
            text_widget.grid(row=0, column=1, sticky="ew", padx=(0, 6))
            row["text_widget"] = text_widget

            if section.get("kind") == "bicolumn":
                buttons = ttk.Frame(row_inner)
                buttons.grid(row=0, column=2, sticky="e", padx=(6, 0))
                self._create_action_button(
                    buttons,
                    text="-",
                    bg="#d9534f",
                    active_bg="#c9302c",
                    command=lambda s=section, i=index: self._delete_ocr_row(s, i),
                ).pack(side="top", anchor="e")
                if self._should_show_ocr_add_button(section, index):
                    self._create_action_button(
                        buttons,
                        text="+",
                        bg="#2ca02c",
                        active_bg="#228b22",
                        command=lambda s=section, i=index + 1: self._insert_ocr_row(s, i),
                    ).pack(side="top", anchor="e", pady=(4, 0))

    def _build_simple_ocr_section(self, parent, title, ocr_text):
        card = ttk.LabelFrame(parent, text=title, padding=10)
        card.pack(fill="x", pady=(0, 12))
        card.columnconfigure(0, weight=1)

        rows_frame = ttk.Frame(card)
        rows_frame.grid(row=0, column=0, sticky="ew")
        section = {
            "kind": "simple",
            "title": title,
            "fixed_type": self._ocr_type_internal(title),
            "fixed_type_display": title,
            "rows_frame": rows_frame,
            "rows": [{"type": self._ocr_type_internal(title), "text": ocr_text}],
            "add_button_label": "+ zone de texte",
        }
        self._render_ocr_section(section)
        return section

    def _build_bicolumn_section(self, parent, title, left_item, right_item):
        card = ttk.LabelFrame(parent, text=title, padding=10)
        card.pack(fill="x", pady=(0, 12))
        card.columnconfigure(0, weight=1)

        section = {
            "kind": "bicolumn",
            "title": title,
            "sides": [],
        }

        left_frame = ttk.Frame(card)
        left_frame.grid(row=0, column=0, sticky="nsew", pady=(0, 12))
        right_frame = ttk.Frame(card)
        right_frame.grid(row=1, column=0, sticky="nsew")

        def populate_side(side_frame, side_label, item):
            side_frame.columnconfigure(0, weight=1)
            ttk.Label(side_frame, text=side_label, font=("Helvetica", 11, "bold")).grid(row=0, column=0, sticky="w", pady=(0, 6))

            rows_frame = ttk.Frame(side_frame)
            rows_frame.grid(row=1, column=0, sticky="nsew")

            blocks = self._ocr_bicolumn_blocks(item["path"])
            rows = [
                {
                    "type": block.get("type", "taxon"),
                    "text": block.get("text", ""),
                }
                for block in blocks
            ]
            rows = self._merge_consecutive_taxon_rows(rows)

            section = {
                "kind": "bicolumn",
                "title": side_label,
                "rows_frame": rows_frame,
                "rows": rows,
                "add_button_label": "+ zone de texte",
            }
            self._render_ocr_section(section)
            return section

        if left_item:
            section["sides"].append(populate_side(left_frame, f"gauche - {left_item['name']}", left_item))
        else:
            ttk.Label(left_frame, text="gauche", font=("Helvetica", 11, "bold")).grid(row=0, column=0, sticky="w")
            section["sides"].append(None)

        if right_item:
            section["sides"].append(populate_side(right_frame, f"droite - {right_item['name']}", right_item))
        else:
            ttk.Label(right_frame, text="droite", font=("Helvetica", 11, "bold")).grid(row=0, column=0, sticky="w")
            section["sides"].append(None)

        return section

    def _build_gallery_ocr_panel(self, parent, exported_items):
        container = ttk.Frame(parent)
        container.grid(row=0, column=0, sticky="nsew")
        container.rowconfigure(0, weight=1)
        container.columnconfigure(0, weight=1)
        container.columnconfigure(1, weight=0)
        sections = []

        scroll_canvas = tk.Canvas(container, highlightthickness=0, bg="#ffffff")
        scrollbar = ttk.Scrollbar(container, orient="vertical", command=scroll_canvas.yview)
        scroll_canvas.configure(yscrollcommand=scrollbar.set)
        scroll_canvas.grid(row=0, column=0, sticky="nsew")
        scrollbar.grid(row=0, column=1, sticky="ns")

        inner_frame = ttk.Frame(scroll_canvas, padding=(0, 0, 8, 0))
        inner_window = scroll_canvas.create_window((0, 0), window=inner_frame, anchor="nw")

        def update_scrollregion(event=None):
            scroll_canvas.configure(scrollregion=scroll_canvas.bbox("all"))

        def resize_inner(event):
            scroll_canvas.itemconfigure(inner_window, width=event.width)

        inner_frame.bind("<Configure>", update_scrollregion)
        scroll_canvas.bind("<Configure>", resize_inner)

        pending_bicolumn_left = None

        for item in exported_items:
            normalized_name = item["name"].lower().replace("-", "_")

            if is_italic_name(normalized_name):
                sections.append(self._build_italic_section(inner_frame, item))
                continue

            if "localisation" in normalized_name:
                ocr_text = self._ocr_image_text(item["path"])
                sections.append(self._build_simple_ocr_section(inner_frame, "localisation", ocr_text))
                continue

            if normalized_name.startswith("genre") or normalized_name.startswith("titre"):
                raw_title_text, fuzzy_title_text = self._ocr_title_text(item["path"])
                sections.append(self._build_simple_ocr_section(inner_frame, "genre", fuzzy_title_text))
                """if fuzzy_title_text != raw_title_text:
                    print("Fuzzy genre:")
                    print(fuzzy_title_text)"""
                continue

            if "bicolonne" in normalized_name or "bi_colonne" in normalized_name:
                if normalized_name.endswith("gauche"):
                    pending_bicolumn_left = item
                    continue

                if normalized_name.endswith("droite"):
                    sections.append(self._build_bicolumn_section(inner_frame, "bicolonne", pending_bicolumn_left, item))
                    pending_bicolumn_left = None
                    continue

                sections.append(self._build_bicolumn_section(inner_frame, "bicolonne", item, None))
                continue

            ocr_text = self._ocr_image_text(item["path"])
            sections.append(self._build_simple_ocr_section(inner_frame, item["name"], ocr_text))

        if pending_bicolumn_left is not None:
            sections.append(self._build_bicolumn_section(inner_frame, "bicolonne", pending_bicolumn_left, None))

        self.current_gallery_sections = sections

        return container

    def _extract_linear_title_blocks(self, gallery_sections):
        blocks = []
        current_block = None
        pending_children = []

        def start_block(title_text):
            nonlocal current_block, pending_children
            current_block = {
                "title": title_text,
                "children": [],
            }
            if pending_children:
                current_block["children"].extend(pending_children)
                pending_children = []
            blocks.append(current_block)

        def append_children(children):
            if current_block is None:
                pending_children.extend(children)
            else:
                current_block["children"].extend(children)

        for section in gallery_sections or []:
            if section is None:
                continue

            kind = section.get("kind")
            if kind == "simple":
                self._sync_ocr_section_rows(section)
                rows = [row for row in section.get("rows", []) if row.get("text", "").strip()]
                section_title = section.get("title")

                if section_title == "localisation":
                    continue

                if section_title in {"titre", "genre"} and rows:
                    title_text = " ".join(rows[0].get("text", "").split())
                    start_block(title_text)
                    continue

                if rows:
                    append_children(
                        [
                            {
                                "type": self._ocr_type_display(section.get("fixed_type", "taxon")),
                                "text": "\n".join(row.get("text", "") for row in rows),
                            }
                        ]
                    )
                continue

            if kind == "bicolumn":
                children = []
                for side in section.get("sides", []):
                    if side is None:
                        continue

                    self._sync_ocr_section_rows(side)
                    for row in side.get("rows", []):
                        row_text = row.get("text", "").strip()
                        if row_text:
                            children.append({"type": self._ocr_type_display(row.get("type", "taxon")), "text": row_text})

                if children:
                    append_children(children)

        if not blocks and pending_children:
            blocks.append({"title": "Titre validé", "children": pending_children})

        return blocks

    def _open_rectangles_gallery(self, exported_items):
        if self.gallery_window is not None and self.gallery_window.winfo_exists():
            self.gallery_window.destroy()

        gallery_window = tk.Toplevel(self)
        self.gallery_window = gallery_window
        gallery_window.title(f"Correction OCR - page {self.current_page + 1}")
        gallery_window.geometry("1600x900")
        gallery_window.columnconfigure(0, weight=1)
        gallery_window.columnconfigure(1, weight=1)
        gallery_window.rowconfigure(0, weight=1)
        gallery_window.rowconfigure(1, weight=0)

        left_panel = ttk.Frame(gallery_window, padding=12)
        left_panel.grid(row=0, column=0, sticky="nsew")
        left_panel.columnconfigure(0, weight=1)
        left_panel.rowconfigure(1, weight=1)

        ttk.Label(left_panel, text="Images générées", font=("Helvetica", 14, "bold")).grid(row=0, column=0, sticky="w", pady=(0, 8))

        scroll_container = ttk.Frame(left_panel)
        scroll_container.grid(row=1, column=0, sticky="nsew")
        scroll_container.columnconfigure(0, weight=1)
        scroll_container.rowconfigure(0, weight=1)

        canvas = tk.Canvas(scroll_container, bg="#f8fafc", highlightthickness=0)
        scrollbar = ttk.Scrollbar(scroll_container, orient="vertical", command=canvas.yview)
        canvas.configure(yscrollcommand=scrollbar.set)
        canvas.grid(row=0, column=0, sticky="nsew")
        scrollbar.grid(row=0, column=1, sticky="ns")

        content_frame = ttk.Frame(canvas)
        content_window = canvas.create_window((0, 0), window=content_frame, anchor="nw")

        def _update_scrollregion(event=None):
            canvas.configure(scrollregion=canvas.bbox("all"))

        def _resize_content(event):
            canvas.itemconfigure(content_window, width=event.width)

        content_frame.bind("<Configure>", _update_scrollregion)
        canvas.bind("<Configure>", _resize_content)

        gallery_window.update_idletasks()
        usable_left_width = min(scroll_container.winfo_width(), canvas.winfo_width(), left_panel.winfo_width())
        left_preview_width = max(int(usable_left_width * 0.72), 320)

        photo_refs = []
        for item in exported_items:
            card = ttk.Frame(content_frame, padding=10, relief="solid")
            card.pack(fill="x", pady=(0, 10))

            ttk.Label(card, text=item["name"], anchor="w", font=("Helvetica", 12, "bold")).pack(anchor="w", pady=(0, 6))

            with Image.open(item["path"]) as image:
                preview = image.convert("RGB")
                normalized_name = item["name"].lower().replace("-", "_")
                if (
                    "bicolonne" in normalized_name
                    or "bi_colonne" in normalized_name
                    or normalized_name.startswith("genre")
                    or normalized_name.startswith("titre")
                    or "localisation" in normalized_name
                ):
                    preview.thumbnail((left_preview_width, 4000), Image.Resampling.LANCZOS)
                else:
                    preview.thumbnail((760, 240), Image.Resampling.LANCZOS)
                photo = ImageTk.PhotoImage(preview)

            photo_refs.append(photo)
            ttk.Label(card, image=photo).pack(anchor="w")

        gallery_window._photo_refs = photo_refs

        right_panel = ttk.Frame(gallery_window, padding=12)
        right_panel.grid(row=0, column=1, sticky="nsew")
        right_panel.columnconfigure(0, weight=1)
        right_panel.rowconfigure(0, weight=1)

        self._refresh_validated_titles_panel()
        self._build_gallery_ocr_panel(right_panel, exported_items)

        bottom_bar = ttk.Frame(gallery_window, padding=(12, 0, 12, 12))
        bottom_bar.grid(row=1, column=0, columnspan=2, sticky="ew")
        bottom_bar.columnconfigure(0, weight=1)
        ttk.Button(
            bottom_bar,
            text="Valider",
            command=lambda: (self._commit_validated_border_page(), gallery_window.destroy()),
        ).grid(row=0, column=1, sticky="e")

    def validate_current_page_borders(self):
        if not self._can_validate_current_page():
            return

        exported_items = self._export_current_page_rectangles()
        if not exported_items:
            return

        self.gallery_export_items = exported_items
        self._open_rectangles_gallery(exported_items)

    def prev_page(self):
        if self.page_count and self.current_page > 0:
            self._show_page(self.current_page - 1)

    def next_page(self):
        if self.page_count and self.current_page < self.page_count - 1:
            self._show_page(self.current_page + 1)

    def go_to_page_from_entry(self, event=None):
        if not self.page_count:
            return

        try:
            page_number = int(self.page_jump_var.get().strip())
        except ValueError:
            self._report_error("Entrez un numero de page valide.")
            self.page_jump_var.set(str(self.current_page + 1))
            return

        if not (1 <= page_number <= self.page_count):
            self._report_error(f"Entrez un numero entre 1 et {self.page_count}.")
            self.page_jump_var.set(str(self.current_page + 1))
            return

        self._show_page(page_number - 1)

    def _fit_image_to_canvas(self):
        if self.current_image is None:
            return None

        canvas_width = max(self.canvas.winfo_width(), 1)
        canvas_height = max(self.canvas.winfo_height(), 1)
        image_width, image_height = self.current_image.size
        scale = min(canvas_width / image_width, canvas_height / image_height)
        scale = min(scale, 1.0) if image_width and image_height else 1.0

        display_width = max(int(image_width * scale), 1)
        display_height = max(int(image_height * scale), 1)
        offset_x = (canvas_width - display_width) // 2
        offset_y = (canvas_height - display_height) // 2

        self.view_scale = scale
        self.offset_x = offset_x
        self.offset_y = offset_y
        self.display_image_size = (display_width, display_height)

        return self.current_image.resize((display_width, display_height), Image.Resampling.LANCZOS)

    def _redraw_canvas(self):
        self.canvas.delete("all")
        self.rect_items.clear()
        self.handle_items.clear()
        self.delete_cross_items.clear()
        self.separator_items.clear()

        if self.current_image is None:
            message = self.canvas_placeholder_text
            if self.current_page not in self.loading_pages:
                message = "Aucune page à afficher"
            self.canvas.create_text(
                self.canvas.winfo_width() // 2,
                self.canvas.winfo_height() // 2,
                text=message,
                fill="white",
                font=("Helvetica", 18),
            )
            return

        display_image = self._fit_image_to_canvas()
        if display_image is None:
            return

        self.current_photo = ImageTk.PhotoImage(display_image)
        self.canvas.create_image(
            self.offset_x,
            self.offset_y,
            anchor="nw",
            image=self.current_photo,
        )

        for rect in self.current_page_data:
            self._draw_rectangle(rect)

    def _image_to_canvas(self, x, y):
        return self.offset_x + x * self.view_scale, self.offset_y + y * self.view_scale

    def _canvas_to_image(self, x, y):
        image_x = (x - self.offset_x) / self.view_scale if self.view_scale else x
        image_y = (y - self.offset_y) / self.view_scale if self.view_scale else y
        if self.current_image is None:
            return image_x, image_y
        image_width, image_height = self.current_image.size
        return (
            clamp(image_x, 0, image_width),
            clamp(image_y, 0, image_height),
        )

    def _draw_rectangle(self, rect):
        if not rect.get("uid"):
            rect["uid"] = self._new_rect_uid()
        else:
            self._next_rect_uid = max(self._next_rect_uid, int(rect["uid"]) + 1)

        canvas_x1, canvas_y1 = self._image_to_canvas(rect["x"], rect["y"])
        canvas_x2, canvas_y2 = self._image_to_canvas(rect["x"] + rect["w"], rect["y"] + rect["h"])
        color = rect_color(rect.get("label", "manual"))
        width = RECT_STROKE_WIDTH

        rect_id = self.canvas.create_rectangle(
            canvas_x1,
            canvas_y1,
            canvas_x2,
            canvas_y2,
            outline=color,
            width=width,
            tags=("rect",),
        )
        self.canvas.tag_bind(rect_id, "<ButtonPress-1>", self._on_canvas_press)
        self.canvas.tag_bind(rect_id, "<B1-Motion>", self._on_canvas_drag)
        self.canvas.tag_bind(rect_id, "<ButtonRelease-1>", self._on_canvas_release)

        rect["canvas_id"] = rect_id
        self.rect_items[rect_id] = rect

        self._draw_separator(rect)
        self._draw_handles(rect)
        self._draw_delete_cross(rect)

    def _get_separator_x(self, rect):
        separator_x = rect.get("center_x")
        if separator_x is None:
            separator_x = rect["x"] + rect["w"] // 2
        return int(clamp(separator_x, rect["x"], rect["x"] + rect["w"]))

    def _draw_separator(self, rect):
        if rect.get("label") != "bicolumn_text":
            return

        separator_x = self._get_separator_x(rect)
        canvas_x = self._image_to_canvas(separator_x, 0)[0]
        x1, y1 = self._image_to_canvas(separator_x, rect["y"])
        x2, y2 = self._image_to_canvas(separator_x, rect["y"] + rect["h"])

        separator_id = self.canvas.create_line(
            canvas_x,
            y1,
            canvas_x,
            y2,
            fill="#111111",
            width=4,
            tags=("separator",),
        )
        self.separator_items[separator_id] = rect["uid"]

    def _find_separator_at_canvas_point(self, canvas_x, canvas_y):
        hits = self.canvas.find_overlapping(canvas_x - 6, canvas_y - 6, canvas_x + 6, canvas_y + 6)
        for item_id in reversed(hits):
            if item_id in self.separator_items:
                return self._get_rect_by_uid(self.separator_items[item_id])
        return None

    def _draw_handles(self, rect):
        if rect.get("uid") != self.active_rect_uid:
            return

        x1, y1 = self._image_to_canvas(rect["x"], rect["y"])
        x2, y2 = self._image_to_canvas(rect["x"] + rect["w"], rect["y"] + rect["h"])
        x_mid = (x1 + x2) / 2
        y_mid = (y1 + y2) / 2
        handle_size = 6

        for hx, hy, name in (
            (x1, y1, "nw"),
            (x_mid, y1, "n"),
            (x2, y1, "ne"),
            (x2, y_mid, "e"),
            (x2, y2, "se"),
            (x_mid, y2, "s"),
            (x1, y2, "sw"),
            (x1, y_mid, "w"),
        ):
            handle = self.canvas.create_rectangle(
                hx - handle_size,
                hy - handle_size,
                hx + handle_size,
                hy + handle_size,
                fill="#ffffff",
                outline="#111111",
                width=1,
                tags=("handle", name),
            )
            self.handle_items[handle] = {"rect_uid": rect["uid"], "handle": name}

    def _draw_delete_cross(self, rect):
        if rect.get("uid") != self.active_rect_uid:
            return

        x2, y_mid = self._image_to_canvas(rect["x"] + rect["w"], rect["y"] + rect["h"] / 2)
        cross_center_x = min(x2 + 14, max(self.canvas.winfo_width() - 14, 14))
        cross_center_y = y_mid
        cross_size = 10

        background = self.canvas.create_oval(
            cross_center_x - cross_size,
            cross_center_y - cross_size,
            cross_center_x + cross_size,
            cross_center_y + cross_size,
            fill="#d9534f",
            outline="#8c2d2a",
            width=1,
            tags=("delete_cross",),
        )
        line1 = self.canvas.create_line(
            cross_center_x - 4,
            cross_center_y - 4,
            cross_center_x + 4,
            cross_center_y + 4,
            fill="white",
            width=2,
            tags=("delete_cross",),
        )
        line2 = self.canvas.create_line(
            cross_center_x - 4,
            cross_center_y + 4,
            cross_center_x + 4,
            cross_center_y - 4,
            fill="white",
            width=2,
            tags=("delete_cross",),
        )

        for item_id in (background, line1, line2):
            self.delete_cross_items[item_id] = rect["uid"]
            self.canvas.tag_bind(item_id, "<ButtonPress-1>", self._on_delete_cross_press)

    def _refresh_active_rectangle_display(self):
        for separator_id, rect_uid in list(self.separator_items.items()):
            if rect_uid == self.active_rect_uid:
                self.canvas.delete(separator_id)
                del self.separator_items[separator_id]

        for handle_id in list(self.handle_items):
            self.canvas.delete(handle_id)
        self.handle_items.clear()

        for cross_id in list(self.delete_cross_items):
            self.canvas.delete(cross_id)
        self.delete_cross_items.clear()

        rect = self._get_rect_by_uid(self.active_rect_uid)
        if rect is None:
            return

        rect_id = rect.get("canvas_id")

        if rect_id is None or rect_id not in self.rect_items:
            self._redraw_canvas()
            return

        for handle_id, handle_data in list(self.handle_items.items()):
            if handle_data["rect_uid"] == rect["uid"]:
                self.canvas.delete(handle_id)
                del self.handle_items[handle_id]

        canvas_x1, canvas_y1 = self._image_to_canvas(rect["x"], rect["y"])
        canvas_x2, canvas_y2 = self._image_to_canvas(rect["x"] + rect["w"], rect["y"] + rect["h"])
        self.canvas.coords(rect_id, canvas_x1, canvas_y1, canvas_x2, canvas_y2)
        self.canvas.itemconfig(
            rect_id,
            outline=rect_color(rect.get("label", "manual")),
            width=RECT_STROKE_WIDTH,
        )

        self._draw_separator(rect)
        self._draw_handles(rect)
        self._draw_delete_cross(rect)

    def _set_active_rectangle(self, rect_uid):
        previous_uid = self.active_rect_uid
        if previous_uid == rect_uid:
            rect = self._get_rect_by_uid(rect_uid)
            self._sync_label_selector_to_rect(rect)
            self._refresh_active_rectangle_display()
            return

        if previous_uid is not None:
            previous_rect = self._get_rect_by_uid(previous_uid)
            if previous_rect is not None:
                previous_canvas_id = previous_rect.get("canvas_id")
                if previous_canvas_id in self.rect_items:
                        self.canvas.itemconfig(previous_canvas_id, width=RECT_STROKE_WIDTH)

        self.active_rect_uid = rect_uid

        if rect_uid is None:
            self._refresh_active_rectangle_display()
            return

        rect = self._get_rect_by_uid(rect_uid)
        self._sync_label_selector_to_rect(rect)
        self._refresh_active_rectangle_display()

    def _sync_label_selector_to_rect(self, rect):
        if rect is None:
            return

        label = label_to_display(label_to_internal(rect.get("label", "manual")))
        if label not in self.label_selector["values"]:
            label = self.label_selector["values"][0]

        if self.new_rect_label.get() != label:
            self.new_rect_label.set(label)

    def _set_rect_label(self, rect_uid, label):
        rect = self._get_rect_by_uid(rect_uid)
        if rect is None:
            return

        label = label_to_internal(label)
        if label_to_display(label) not in self.label_selector["values"]:
            label = label_to_internal(self.label_selector["values"][0])

        rect["label"] = label
        self._save_current_page_state()
        if rect_uid == self.active_rect_uid:
            self._refresh_active_rectangle_display()
        else:
            self._redraw_canvas()

    def _on_label_selector_change(self, event=None):
        label = label_to_internal(self.new_rect_label.get())

        if self.active_rect_uid is not None:
            self._set_rect_label(self.active_rect_uid, label)
            return

        if self.new_rect_label.get() not in self.label_selector["values"]:
            self.new_rect_label.set(self.label_selector["values"][0])

    def _get_rect_by_uid(self, rect_uid):
        for rect in self.current_page_data:
            if rect.get("uid") == rect_uid:
                return rect
        return None

    def _find_rect_at_canvas_point(self, canvas_x, canvas_y):
        image_x, image_y = self._canvas_to_image(canvas_x, canvas_y)

        for rect in reversed(self.current_page_data):
            if (
                rect["x"] <= image_x <= rect["x"] + rect["w"] and
                rect["y"] <= image_y <= rect["y"] + rect["h"]
            ):
                return rect
        return None

    def _find_handle_at_canvas_point(self, canvas_x, canvas_y):
        hits = self.canvas.find_overlapping(canvas_x, canvas_y, canvas_x, canvas_y)
        for item_id in reversed(hits):
            if item_id in self.handle_items:
                return self.handle_items[item_id]
        return None

    def _find_delete_cross_at_canvas_point(self, canvas_x, canvas_y):
        hits = self.canvas.find_overlapping(canvas_x, canvas_y, canvas_x, canvas_y)
        for item_id in reversed(hits):
            if item_id in self.delete_cross_items:
                return self.delete_cross_items[item_id]
        return None

    def _on_delete_cross_press(self, event):
        delete_cross_uid = self._find_delete_cross_at_canvas_point(event.x, event.y)
        if delete_cross_uid is None:
            return "break"

        self._set_active_rectangle(delete_cross_uid)
        self._delete_selected_rectangle()
        return "break"

    def _clear_interaction_state(self, keep_selection=False):
        self.action_mode = None
        self.drag_start_image = None
        self.drag_anchor_rect = None
        if not keep_selection:
            self.active_rect_uid = None
        self.temp_rect_id = None

    def _on_canvas_resize(self, event):
        if self.current_image is not None:
            self._redraw_canvas()

    def start_add_rectangle(self):
        if self.current_image is None:
            return
        self._clear_interaction_state()
        self.action_mode = "add"

    def _cancel_action(self, event=None):
        if self.action_mode == "add" and self.temp_rect_id is not None:
            self.canvas.delete(self.temp_rect_id)
            self.temp_rect_id = None
        self._clear_interaction_state()
        self._redraw_canvas()

    def _on_canvas_press(self, event):
        if self.current_image is None:
            return

        canvas_x, canvas_y = event.x, event.y
        image_x, image_y = self._canvas_to_image(canvas_x, canvas_y)

        if self.action_mode == "add":
            self.drag_start_image = (image_x, image_y)
            self.temp_rect_id = self.canvas.create_rectangle(
                canvas_x,
                canvas_y,
                canvas_x,
                canvas_y,
                outline=rect_color(label_to_internal(self.new_rect_label.get())),
                width=RECT_STROKE_WIDTH,
                dash=(4, 2),
            )
            return

        delete_cross_uid = self._find_delete_cross_at_canvas_point(canvas_x, canvas_y)
        if delete_cross_uid is not None:
            self._set_active_rectangle(delete_cross_uid)
            self._delete_selected_rectangle()
            return

        handle = self._find_handle_at_canvas_point(canvas_x, canvas_y)
        if handle:
            self._set_active_rectangle(handle["rect_uid"])
            self.action_mode = f"resize:{handle['handle']}"
            self.drag_start_image = (image_x, image_y)
            self.drag_anchor_rect = self._get_rect_by_uid(self.active_rect_uid).copy()
            return

        separator_rect = self._find_separator_at_canvas_point(canvas_x, canvas_y)
        if separator_rect:
            self._set_active_rectangle(separator_rect["uid"])
            self.action_mode = "separator"
            self.drag_start_image = (image_x, image_y)
            self.drag_anchor_rect = separator_rect.copy()
            return

        rect = self._find_rect_at_canvas_point(canvas_x, canvas_y)
        if rect:
            self._set_active_rectangle(rect["uid"])
            self.action_mode = "move"
            self.drag_start_image = (image_x, image_y)
            self.drag_anchor_rect = rect.copy()
            return

        self._clear_interaction_state()
        self._redraw_canvas()

    def _on_canvas_drag(self, event):
        if self.current_image is None or self.action_mode is None:
            return

        canvas_x, canvas_y = event.x, event.y
        image_x, image_y = self._canvas_to_image(canvas_x, canvas_y)

        if self.action_mode == "add" and self.temp_rect_id is not None and self.drag_start_image:
            x0, y0 = self.drag_start_image
            x1 = clamp(min(x0, image_x), 0, self.current_image.size[0])
            y1 = clamp(min(y0, image_y), 0, self.current_image.size[1])
            x2 = clamp(max(x0, image_x), 0, self.current_image.size[0])
            y2 = clamp(max(y0, image_y), 0, self.current_image.size[1])

            self.canvas.coords(
                self.temp_rect_id,
                *self._image_to_canvas(x1, y1),
                *self._image_to_canvas(x2, y2),
            )
            return

        rect = self._get_rect_by_uid(self.active_rect_uid)
        if rect is None or self.drag_anchor_rect is None or self.drag_start_image is None:
            return

        dx = image_x - self.drag_start_image[0]
        dy = image_y - self.drag_start_image[1]
        base = self.drag_anchor_rect

        if self.action_mode == "move":
            rect["x"] = int(clamp(base["x"] + dx, 0, self.current_image.size[0] - 1))
            rect["y"] = int(clamp(base["y"] + dy, 0, self.current_image.size[1] - 1))
            if rect.get("label") == "bicolumn_text":
                separator_x = base.get("center_x", base["x"] + base["w"] // 2) + dx
                rect["center_x"] = int(clamp(separator_x, rect["x"], rect["x"] + rect["w"]))
        elif self.action_mode == "separator":
            if rect.get("label") == "bicolumn_text":
                separator_x = base.get("center_x", base["x"] + base["w"] // 2) + dx
                rect["center_x"] = int(clamp(separator_x, rect["x"], rect["x"] + rect["w"]))
        elif self.action_mode.startswith("resize:"):
            handle = self.action_mode.split(":", 1)[1]
            x1 = base["x"]
            y1 = base["y"]
            x2 = base["x"] + base["w"]
            y2 = base["y"] + base["h"]

            if "w" in handle:
                x1 = clamp(image_x, 0, x2 - 1)
            if "e" in handle:
                x2 = clamp(image_x, x1 + 1, self.current_image.size[0])
            if "n" in handle:
                y1 = clamp(image_y, 0, y2 - 1)
            if "s" in handle:
                y2 = clamp(image_y, y1 + 1, self.current_image.size[1])

            rect["x"] = int(x1)
            rect["y"] = int(y1)
            rect["w"] = max(1, int(x2 - x1))
            rect["h"] = max(1, int(y2 - y1))
            if rect.get("label") == "bicolumn_text":
                separator_x = base.get("center_x", base["x"] + base["w"] // 2)
                rect["center_x"] = int(clamp(separator_x, rect["x"], rect["x"] + rect["w"]))

        self._refresh_active_rectangle_display()

    def _on_canvas_release(self, event):
        if self.current_image is None or self.action_mode is None:
            return

        if self.action_mode == "add" and self.temp_rect_id is not None and self.drag_start_image:
            start_x, start_y = self.drag_start_image
            end_x, end_y = self._canvas_to_image(event.x, event.y)

            x1 = int(clamp(min(start_x, end_x), 0, self.current_image.size[0]))
            y1 = int(clamp(min(start_y, end_y), 0, self.current_image.size[1]))
            x2 = int(clamp(max(start_x, end_x), 0, self.current_image.size[0]))
            y2 = int(clamp(max(start_y, end_y), 0, self.current_image.size[1]))

            self.canvas.delete(self.temp_rect_id)
            self.temp_rect_id = None

            if (x2 - x1) > 4 and (y2 - y1) > 4:
                new_rect = {
                    "uid": self._new_rect_uid(),
                    "x": x1,
                    "y": y1,
                    "w": x2 - x1,
                    "h": y2 - y1,
                    "label": label_to_internal(self.new_rect_label.get()),
                }
                if new_rect["label"] == "bicolumn_text":
                    new_rect["center_x"] = x1 + (x2 - x1) // 2
                self.current_page_data.append(new_rect)
                self._set_active_rectangle(new_rect["uid"])

            self._clear_interaction_state(keep_selection=True)
            self._redraw_canvas()
            return

        self._clear_interaction_state(keep_selection=True)
        self._refresh_active_rectangle_display()

    def _delete_selected_rectangle(self, event=None):
        if self.active_rect_uid is None:
            return

        rect = self._get_rect_by_uid(self.active_rect_uid)
        if rect in self.current_page_data:
            self.current_page_data.remove(rect)

        self._save_current_page_state()
        self._clear_interaction_state()
        self._redraw_canvas()

    def _build_rectangle_payload(self):
        self._save_current_page_state()

        return {
            "pdf_path": self.pdf_path,
            "pages": [
                {
                    "page_index": index,
                    "page_number": index + 1,
                    "rectangles": [
                        {
                            "x": rect["x"],
                            "y": rect["y"],
                            "w": rect["w"],
                            "h": rect["h"],
                            "label": rect.get("label", "manual"),
                            **({"center_x": rect["center_x"]} if rect.get("label") == "bicolumn_text" and "center_x" in rect else {}),
                        }
                        for rect in (page_rects or [])
                    ],
                }
                for index, page_rects in enumerate(self.page_annotations)
            ],
        }
    
    def _report_error(self, message):
        messagebox.showerror("Erreur", message)

if __name__ == "__main__":
    app = PinaxBauhinDataApp()
    app.mainloop()